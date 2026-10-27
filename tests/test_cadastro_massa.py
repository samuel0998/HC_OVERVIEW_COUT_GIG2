"""Cadastro em massa por planilha (Novo HC): template com aba de Instruções
em primeiro lugar, e upload que valida tudo antes de gravar (tudo-ou-nada)."""
import io
import unittest
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
from flask import Flask

from models import db
from models.hc_gig2 import HCGig2
from models.registro_atividade import RegistroAtividade
from routes.hc import hc_bp

CABECALHO = ["Nome completo", "Login", "Cargo", "Área", "Turno", "Agência"]


class CadastroMassaTest(unittest.TestCase):
    def setUp(self):
        app = Flask(__name__)
        app.config.update(TESTING=True, LOGIN_DISABLED=True, SQLALCHEMY_DATABASE_URI="sqlite:///:memory:")
        db.init_app(app)
        app.register_blueprint(hc_bp)
        self.context = app.app_context()
        self.context.push()
        self.addCleanup(self.context.pop)
        self.addCleanup(db.session.remove)
        HCGig2.__table__.create(db.engine)
        RegistroAtividade.__table__.create(db.engine)
        self.client = app.test_client()
        user = patch("routes.hc.current_user", SimpleNamespace(
            is_authenticated=True, login="tester", nome="Testador", can_add_colaborador=True, can_edit=True,
        ))
        user.start()
        self.addCleanup(user.stop)

    def _planilha(self, linhas, aba="Cadastro", cabecalho=CABECALHO):
        wb = openpyxl.Workbook()
        wb.active.title = "Instruções"
        ws = wb.create_sheet(aba)
        ws.append(cabecalho)
        for linha in linhas:
            ws.append(linha)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _enviar(self, buf):
        return self.client.post(
            "/api/hc/cadastro-massa",
            data={"arquivo": (buf, "cadastro.xlsx")},
            content_type="multipart/form-data",
        )

    def test_template_tem_instrucoes_como_primeira_aba(self):
        resposta = self.client.get("/api/hc/cadastro-massa/template")
        self.assertEqual(resposta.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resposta.data))
        self.assertEqual(wb.sheetnames, ["Instruções", "Cadastro", "Listas"])
        cabecalho = [c.value for c in wb["Cadastro"][1]]
        self.assertEqual(cabecalho, CABECALHO)
        # Listas suspensas de Cargo/Área/Turno na aba Cadastro.
        self.assertEqual(len(wb["Cadastro"].data_validations.dataValidation), 3)

    def test_cadastra_todas_as_linhas_validas(self):
        resposta = self._enviar(self._planilha([
            ["Maria Souza", "msouza", "Associado", "INBOUND", "BLUE DAY", "ID LOGISTICS"],
            ["João Lima", "jlima", "pit trainee", "outbound", "red night", None],
            ["Sem Login", None, "PIT", None, None, None],
        ]))
        self.assertEqual(resposta.status_code, 201, resposta.get_json())
        self.assertEqual(resposta.get_json()["cadastrados"], 3)

        maria = HCGig2.query.filter_by(login="msouza").first()
        self.assertEqual(maria.status, "Treinamento")
        self.assertEqual(maria.area, "INBOUND")
        self.assertEqual(maria.turno, "BLUE DAY")
        self.assertEqual(maria.agencia, "ID LOGISTICS")
        self.assertIsNotNone(maria.treinamento_inicio_em)

        joao = HCGig2.query.filter_by(login="jlima").first()
        self.assertEqual(joao.cargo, "PIT")          # grava e conta como PIT
        self.assertTrue(joao.pit_trainee)            # rótulo preservado
        self.assertEqual(joao.area, "OUTBOUND")      # normalizado pra lista oficial
        self.assertEqual(joao.turno, "RED NIGHT")

        self.assertEqual(RegistroAtividade.query.filter_by(tipo="adicao").count(), 3)

    def test_qualquer_erro_nao_cadastra_nada(self):
        resposta = self._enviar(self._planilha([
            ["Válido", "valido1", "Associado", "INBOUND", "BLUE DAY", None],
            ["Cargo Errado", "x2", "Gerentão", "INBOUND", "BLUE DAY", None],
            ["Turno Errado", "x3", "Associado", "INBOUND", "MADRUGADA", None],
            [None, "x4", "Associado", None, None, None],
        ]))
        self.assertEqual(resposta.status_code, 400)
        data = resposta.get_json()
        linhas_com_erro = {e["linha"] for e in data["erros"]}
        self.assertEqual(linhas_com_erro, {3, 4, 5})
        self.assertEqual(HCGig2.query.count(), 0)  # tudo-ou-nada

    def test_login_repetido_na_planilha_ou_no_sistema(self):
        db.session.add(HCGig2(nome_completo="Já Existe", login="existente", cargo="Associado", status="OPERACIONAL"))
        db.session.commit()
        resposta = self._enviar(self._planilha([
            ["A", "dup", "Associado", None, None, None],
            ["B", "DUP", "Associado", None, None, None],
            ["C", "existente", "Associado", None, None, None],
        ]))
        self.assertEqual(resposta.status_code, 400)
        erros = {e["linha"]: " ".join(e["erros"]) for e in resposta.get_json()["erros"]}
        self.assertIn("repetido", erros[3])
        self.assertIn("já cadastrado", erros[4])
        self.assertEqual(HCGig2.query.count(), 1)

    def test_linhas_vazias_sao_ignoradas(self):
        resposta = self._enviar(self._planilha([
            ["Um", None, "Associado", None, None, None],
            [None, None, None, None, None, None],
            ["Dois", None, "PIT", None, None, None],
        ]))
        self.assertEqual(resposta.status_code, 201, resposta.get_json())
        self.assertEqual(HCGig2.query.count(), 2)

    def test_aba_cadastro_renomeada_e_rejeitada(self):
        resposta = self._enviar(self._planilha([["Um", None, "Associado", None, None, None]], aba="Planilha1"))
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("Cadastro", resposta.get_json()["erro"])

    def test_planilha_sem_linhas(self):
        resposta = self._enviar(self._planilha([]))
        self.assertEqual(resposta.status_code, 400)

    def test_exige_permissao_de_cadastro(self):
        with patch("routes.hc.current_user", SimpleNamespace(is_authenticated=True, login="x", nome="X", can_add_colaborador=False)):
            resposta = self._enviar(self._planilha([["Um", None, "Associado", None, None, None]]))
        self.assertEqual(resposta.status_code, 403)


if __name__ == "__main__":
    unittest.main()
