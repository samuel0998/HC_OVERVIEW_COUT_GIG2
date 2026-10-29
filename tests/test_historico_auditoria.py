"""Aba Auditoria de Histórico: métricas por tipo/dia, período em horário de
São Paulo, comparativo com período anterior, validações e relatório gerencial
(ver routes.hc.auditoria_historico / exportar_auditoria_historico)."""
import io
import json
import unittest
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
from flask import Flask

from models import db
from models.hc_gig2 import HCGig2
from models.registro_atividade import RegistroAtividade
from routes.hc import hc_bp


class HistoricoAuditoriaTest(unittest.TestCase):
    def setUp(self):
        app = Flask(__name__)
        app.config.update(TESTING=True, LOGIN_DISABLED=True, SQLALCHEMY_DATABASE_URI="sqlite:///:memory:")
        db.init_app(app)
        app.register_blueprint(hc_bp)
        self.context = app.app_context()
        self.context.push()
        self.addCleanup(self.context.pop)
        self.addCleanup(db.session.remove)
        HCGig2.__table__.create(db.engine)
        RegistroAtividade.__table__.create(db.engine)
        self.client = app.test_client()
        user = patch("routes.hc.current_user", SimpleNamespace(
            is_authenticated=True, login="tester", nome="Testador", can_historico=True,
        ))
        user.start()
        self.addCleanup(user.stop)

    def _colaborador(self, **kwargs):
        kwargs.setdefault("cargo", "Associado")
        kwargs.setdefault("status", "OPERACIONAL")
        c = HCGig2(**kwargs)
        db.session.add(c)
        db.session.commit()
        return c

    def _registro(self, op, tipo, ts_utc, usuario="tester", descricao=None, ant=None, nov=None):
        reg = RegistroAtividade(
            tipo=tipo,
            operador_id=op.id if op else None,
            operador_login=op.login if op else None,
            operador_nome=op.nome_completo if op else None,
            usuario_login=usuario,
            usuario_nome=usuario,
            descricao=descricao or f"{tipo} {ts_utc.isoformat()}",
            dados_anteriores=ant,
            dados_novos=nov,
            timestamp=ts_utc,
        )
        db.session.add(reg)
        db.session.commit()
        return reg

    def _auditoria(self, **params):
        resposta = self.client.get("/api/hc/historico/auditoria", query_string=params)
        self.assertEqual(resposta.status_code, 200, resposta.get_json())
        return resposta.get_json()

    def test_quantidade_por_tipo_e_metricas(self):
        a = self._colaborador(nome_completo="Ana", login="ana", area="INBOUND")
        b = self._colaborador(nome_completo="Bia", login="bia", area="OUTBOUND")
        self._registro(a, "adicao", datetime(2026, 6, 10, 15))
        self._registro(a, "edicao", datetime(2026, 6, 11, 15))
        self._registro(b, "edicao", datetime(2026, 6, 11, 16), usuario="outro")
        self._registro(b, "edicao_status", datetime(2026, 6, 12, 15), usuario="sistema")

        d = self._auditoria(data_de="2026-06-10", data_ate="2026-06-12")
        por_tipo = {t["tipo"]: t["total"] for t in d["por_tipo"]}
        self.assertEqual(por_tipo, {"edicao": 2, "adicao": 1, "edicao_status": 1})
        self.assertEqual(d["por_tipo"][0]["tipo"], "edicao")  # ordenado por quantidade

        m = d["metricas"]
        self.assertEqual(m["total"], 4)
        self.assertEqual(m["colaboradores"], 2)
        self.assertEqual(m["usuarios"], 2)  # "sistema" não conta como usuário
        self.assertEqual(m["automaticos"], 1)
        self.assertEqual(m["pico_dia"], "2026-06-11")
        self.assertEqual(m["pico_total"], 2)
        self.assertEqual(d["por_dia"]["2026-06-11"]["por_tipo"]["edicao"], 2)
        self.assertEqual({s["nome"]: s["total"] for s in d["por_setor"]}, {"INBOUND": 2, "OUTBOUND": 2})

    def test_dia_e_hora_usam_horario_de_sao_paulo(self):
        a = self._colaborador(nome_completo="Ana", login="ana")
        # 01:30 UTC do dia 11 = 22:30 do dia 10 em São Paulo.
        self._registro(a, "edicao", datetime(2026, 6, 11, 1, 30))
        d = self._auditoria(data_de="2026-06-10", data_ate="2026-06-10")
        self.assertEqual(d["metricas"]["total"], 1)
        self.assertIn("2026-06-10", d["por_dia"])
        self.assertEqual(d["por_hora"][22], 1)

        vazio = self._auditoria(data_de="2026-06-11", data_ate="2026-06-11")
        self.assertEqual(vazio["metricas"]["total"], 0)

    def test_comparativo_com_periodo_anterior(self):
        a = self._colaborador(nome_completo="Ana", login="ana")
        self._registro(a, "edicao", datetime(2026, 6, 5, 15))       # período anterior
        self._registro(a, "edicao", datetime(2026, 6, 6, 15))       # período anterior
        for dia in (12, 13, 14):
            self._registro(a, "edicao", datetime(2026, 6, dia, 15))
        d = self._auditoria(data_de="2026-06-10", data_ate="2026-06-16")
        self.assertEqual(d["periodo_anterior"], {"data_de": "2026-06-03", "data_ate": "2026-06-09", "total": 2})
        self.assertEqual(d["metricas"]["variacao_pct"], 50.0)

    def test_periodo_padrao_e_ultimos_30_dias(self):
        d = self._auditoria()
        self.assertEqual(d["periodo"]["dias"], 30)

    def test_filtro_por_tipo_e_setor_atual(self):
        a = self._colaborador(nome_completo="Ana", login="ana", area="INBOUND")
        b = self._colaborador(nome_completo="Bia", login="bia", area="OUTBOUND")
        self._registro(a, "edicao", datetime(2026, 6, 10, 15))
        self._registro(a, "adicao", datetime(2026, 6, 10, 15))
        self._registro(b, "edicao", datetime(2026, 6, 10, 15))
        d = self._auditoria(data_de="2026-06-10", data_ate="2026-06-10", tipo="edicao", area="INBOUND")
        self.assertEqual(d["metricas"]["total"], 1)

    def _validacao(self, dados, chave):
        return next(v for v in dados["validacoes"] if v["chave"] == chave)

    def test_validacoes(self):
        a = self._colaborador(nome_completo="Ana", login="ana")
        ts = datetime(2026, 6, 10, 15)
        # duplicado: mesmo evento repetido em 1 minuto
        self._registro(a, "edicao", ts, descricao="Mudou setor")
        self._registro(a, "edicao", ts + timedelta(minutes=1), descricao="Mudou setor")
        # sem usuário responsável
        self._registro(a, "edicao", ts + timedelta(hours=1), usuario="")
        # edição sem alteração
        mesmo = json.dumps({"area": "INBOUND"})
        self._registro(a, "edicao", ts + timedelta(hours=2), ant=mesmo, nov=mesmo)
        # órfão: colaborador apagado sem registro de saída
        sumido = self._colaborador(nome_completo="Sumido", login="sumido")
        self._registro(sumido, "edicao", ts + timedelta(hours=3))
        # exclusão aponta pra id inexistente de propósito: não é órfão
        excluido = self._colaborador(nome_completo="Excluido", login="excl")
        self._registro(excluido, "exclusao", ts + timedelta(hours=4))
        db.session.delete(sumido)
        db.session.delete(excluido)
        db.session.commit()

        d = self._auditoria(data_de="2026-06-10", data_ate="2026-06-10")
        self.assertEqual(self._validacao(d, "duplicados")["total"], 1)
        self.assertEqual(self._validacao(d, "sem_usuario")["total"], 1)
        self.assertEqual(self._validacao(d, "sem_alteracao")["total"], 1)
        orfaos = self._validacao(d, "orfaos")
        self.assertEqual(orfaos["total"], 1)
        self.assertEqual(orfaos["itens"][0]["login"], "sumido")
        self.assertEqual(self._validacao(d, "sem_usuario")["severidade"], "critico")
        self.assertEqual(self._validacao(d, "sem_colaborador")["severidade"], "ok")
        self.assertGreaterEqual(d["metricas"]["inconsistencias"], 3)

    def test_validacao_ls_vencido(self):
        self._colaborador(nome_completo="Em LS", login="ls1", status="LS", area="OUTBOUND", ls_area_origem="INBOUND",
                          ls_retorno_em=datetime.utcnow() - timedelta(hours=2))
        self._colaborador(nome_completo="LS ok", login="ls2", status="LS",
                          ls_retorno_em=datetime.utcnow() + timedelta(days=1))
        d = self._auditoria()
        ls = self._validacao(d, "ls_vencidos")
        self.assertEqual(ls["total"], 1)
        self.assertEqual(ls["itens"][0]["login"], "ls1")

    def test_relatorio_gerencial_excel(self):
        a = self._colaborador(nome_completo="Ana", login="ana", area="INBOUND")
        self._registro(a, "adicao", datetime(2026, 6, 10, 15))
        self._registro(a, "edicao", datetime(2026, 6, 11, 15))
        resposta = self.client.get("/api/hc/historico/auditoria/export?data_de=2026-06-10&data_ate=2026-06-12")
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("spreadsheetml", resposta.content_type)
        wb = openpyxl.load_workbook(io.BytesIO(resposta.data))
        self.assertEqual(wb.sheetnames, [
            "Resumo", "Por tipo", "Por dia", "Por usuário", "Setor-Cargo-Turno",
            "Validações", "Validações - detalhe", "Detalhe",
        ])
        por_dia = list(wb["Por dia"].iter_rows(values_only=True))
        self.assertEqual(len(por_dia) - 1, 3)  # um por dia do período, inclusive o vazio
        self.assertEqual(por_dia[1][0], "10/06/2026")
        self.assertEqual(len(list(wb["Detalhe"].iter_rows())) - 1, 2)

    def test_exige_permissao_de_historico(self):
        with patch("routes.hc.current_user", SimpleNamespace(is_authenticated=True, login="x", nome="X", can_historico=False)):
            self.assertEqual(self.client.get("/api/hc/historico/auditoria").status_code, 403)
            self.assertEqual(self.client.get("/api/hc/historico/auditoria/export").status_code, 403)


if __name__ == "__main__":
    unittest.main()
